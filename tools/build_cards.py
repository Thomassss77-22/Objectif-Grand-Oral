#!/usr/bin/env python3
"""Compile les documents du dossier « Grand Oral/ » en paquets de cartes pour l'application.

Sources lues :
  - Fiches pratiques/OB_Partie_*.docx : les 56 fiches du cours (fiche-mémoire + dossier) ;
  - Annexes/CARTE DU COURS … .docx : intitulés des parties et des fiches ;
  - Q&A/GrandOral_Banque_questions_jury.pdf : 580 questions du jury (oct. 2026) ;
  - PréBarreau/500 Q° 2025.pdf : 500 questions-réponses (éd. 2025) ;
  - Q&A/Grand Oral - MAJ Q&R personnalites et dates (oct. 2026).docx : mises à jour des réponses datées.

Sortie : app/data/cours.json, app/data/jury.json, app/data/q500.json.

Dépendances : python-docx, pdfplumber, pdftotext (poppler-utils), pandoc.
Usage : python3 tools/build_cards.py
"""
import glob
import hashlib
import html
import json
import os
import re
import subprocess
import sys

import docx
import pdfplumber

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "Grand Oral")
OUT = os.path.join(ROOT, "app", "data")

# Code couleur des fiches : textes (bleu), jurisprudence (rose), doctrine (vert).
COLORS = {"1E88E5": "t", "1E88C7": "t", "E91E63": "j", "E91D62": "j", "2E7D32": "d"}


def cid(prefix, text):
    return f"{prefix}-{hashlib.sha1(text.encode()).hexdigest()[:8]}"


def clean(s):
    return re.sub(r"\s+", " ", s.replace(" ", " ")).strip()


# ---------- Paragraphes Word → segments stylés → HTML ----------
def segments(p):
    """Liste de (texte, style) où style ∈ {"", "b", "t", "j", "d"} (gras = à mémoriser)."""
    out = []
    for r in p.runs:
        if not r.text:
            continue
        c = r.font.color.rgb if r.font.color is not None and r.font.color.type is not None else None
        st = COLORS.get(str(c), "")
        if not st and r.bold:
            st = "b"
        if out and out[-1][1] == st:
            out[-1] = (out[-1][0] + r.text, st)
        else:
            out.append((r.text, st))
    return out


def to_html(segs):
    parts = []
    for text, st in segs:
        t = html.escape(text.replace(" ", " "), quote=False)
        if st == "b":
            parts.append(f"<b>{t}</b>")
        elif st:
            parts.append(f'<i class="{st}">{t}</i>')
        else:
            parts.append(t)
    return re.sub(r"\s+", " ", "".join(parts)).strip()


def plain(segs):
    return clean("".join(t for t, _ in segs))


def split_at(segs, pos, sep_len):
    """Coupe les segments à la position pos du texte brut ; retire le séparateur."""
    left, right, n = [], [], 0
    for text, st in segs:
        a, b = n, n + len(text)
        if b <= pos:
            left.append((text, st))
        elif a >= pos + sep_len:
            right.append((text, st))
        else:
            if pos > a:
                left.append((text[: pos - a], st))
            if b > pos + sep_len:
                right.append((text[max(0, pos + sep_len - a):], st))
        n = b
    return left, right


def lead_split(segs, seps=(" : ", " — ")):
    """Sépare « intitulé : contenu » au premier séparateur trouvé (le plus à gauche)."""
    raw = "".join(t for t, _ in segs)
    best = None
    for s in seps:
        i = raw.find(s)
        if i > 0 and (best is None or i < best[0]):
            best = (i, len(s))
    if not best:
        return None
    return split_at(segs, *best)


def strip_star(segs):
    if segs and segs[0][0].lstrip().startswith("★"):
        t, st = segs[0]
        t = t.lstrip()[1:].lstrip()
        segs = ([(t, st)] if t else []) + segs[1:]
    return segs


# ---------- Cours : intitulés ----------
def course_map():
    path = glob.glob(os.path.join(SRC, "Annexes", "CARTE DU COURS*.docx"))[0]
    d = docx.Document(path)
    parts, fiches, cur = [], {}, None
    for p in d.paragraphs:
        t = clean(p.text)
        m = re.match(r"^PARTIE ([IVX]+) — (.+)$", t)
        if m and not fiches.get("_done"):
            cur = {"id": m.group(1), "title": m.group(2).capitalize(), "fiches": []}
            if not any(x["id"] == cur["id"] for x in parts):
                parts.append(cur)
            continue
        m = re.match(r"^Fiche (\d+) — (.+?) p\. \d+$", t)
        if m and cur:
            n = int(m.group(1))
            if n not in fiches:
                fiches[n] = m.group(2)
                cur["fiches"].append(n)
    # Intitulés lisibles : « Éléments introductifs » plutôt que « ÉLÉMENTS INTRODUCTIFS ».
    for x in parts:
        x["title"] = x["title"][0].upper() + x["title"][1:].lower()
        x["title"] = x["title"].replace("l'union", "l'Union").replace("l’union", "l’Union")
    return parts, fiches


# ---------- Cours : fiches ----------
def parse_fiches(titles):
    fiches = {}
    files = sorted(glob.glob(os.path.join(SRC, "Fiches pratiques", "OB_Partie_*.docx")))
    for f in files:
        d = docx.Document(f)
        cur = None
        h2 = h3 = ""
        for p in d.paragraphs:
            style = p.style.name if p.style is not None else ""
            text = clean(p.text)
            if not text:
                continue
            if style == "Heading 1":
                m = re.match(r"^FICHE N° (\d+)", text)
                cur = None
                if m:
                    n = int(m.group(1))
                    cur = fiches.setdefault(n, {"n": n, "title": titles.get(n, text), "memo": {}, "sujets": [], "items": []})
                h2 = h3 = ""
                continue
            if cur is None:
                continue
            if style == "Heading 2":
                h2, h3 = text, ""
                continue
            if style == "Heading 3":
                h3 = text
                continue
            if style != "List Paragraph":
                continue
            cur["items"].append((h2, h3, segments(p)))
    return fiches


MEMO_KEYS = {
    "Accroche": "accroche",
    "Définition en trois lignes": "definition",
    "Les cinq références incontournables": "refs",
    "Deux matrices de plan": "plans",
    "Trois pièges": "pieges",
}

# Décision de justice en tête d'un paragraphe : sert à éviter les doublons entre la fiche-mémoire et le dossier.
MONTHS = "janvier|février|mars|avril|mai|juin|juillet|août|septembre|octobre|novembre|décembre"


def decision_key(s):
    s = s.replace("1er", "1")
    m = re.search(r"(\d{1,2})\s+(" + MONTHS + r")\s+(\d{4})", s)
    court = re.match(r"\s*([A-Za-zÀ-ÿ]+)", s)
    if m:
        return (court.group(1).lower() if court else "", m.group(1), m.group(2), m.group(3))
    m = re.search(r"n°\s*([\w\-/.]+)", s)
    return ("n", m.group(1)) if m else None


def plan_html(h):
    # Un plan par ligne : « I. … / II. … » et sous-parties « (A. … ; B. …) ».
    return re.sub(r"\s/\s(?=II\.)", "<br>", h)


def build_fiche(fi):
    n, title = fi["n"], fi["title"]
    memo = {k: [] for k in MEMO_KEYS.values()}
    cards, seen_dec, seen_terms = [], set(), set()
    ctx_title = f"Fiche {n} · {title}"

    # core : cartes de la fiche-mémoire et du dispositif « Grand Oral » (le socle) ; les autres
    # viennent du dossier détaillé.
    def add(kind, q, a, ctx="", core=False):
        if not a or len(plain_text(a)) < 8:
            return
        c = {"id": cid(f"f{n}", kind + q), "k": kind, "q": q, "a": a, "ctx": ctx}
        if core:
            c["core"] = 1
        cards.append(c)

    for h2, h3, segs in fi["items"]:
        raw = plain(segs)
        memo_part = h2.startswith("◆")
        if memo_part:
            key = MEMO_KEYS.get(h3)
            if not key:
                continue
            memo[key].append(to_html(strip_star(segs)))
            if key == "definition":
                sp = lead_split(segs, (" : ",))
                if sp and len(plain(sp[0])) <= 70:
                    term = plain(sp[0])
                    seen_terms.add(term.lower())
                    add("def", ask_term(term), to_html(sp[1]), ctx_title, True)
            elif key == "refs":
                s = strip_star(segs)
                sp = lead_split(s, (" : ", " — "))
                if sp:
                    ref = plain(sp[0])
                    dk = decision_key(ref)
                    if dk:
                        seen_dec.add(dk)
                    add("ref", f"Portée de : {ref}", to_html(sp[1]), ctx_title, True)
            elif key == "plans":
                m = re.match(r"^Sujet\s+(«.+?»)\s*[—:-]\s*", raw)
                if m:
                    sp = split_at(segs, m.end(), 0)
                    add("plan", f"Plan possible pour le sujet {m.group(1)}", plan_html(to_html(sp[1])), ctx_title, True)
            continue

        sub = h3 or h2
        if "Sujets d’exposé" in sub or "Sujets d'exposé" in sub:
            fi["sujets"].append(to_html(segs))
            continue
        if "Angles morts" in sub:
            body = re.sub(r"^Q\.\s*", "", raw)
            if body == raw:
                continue
            m = re.search(r"\?\s+—\s+", raw) or re.search(r"\s—\s", raw)
            if not m:
                continue
            q = clean(re.sub(r"^Q\.\s*", "", raw[: m.start() + (1 if raw[m.start()] == "?" else 0)]))
            a = split_at(segs, m.end(), 0)[1]
            add("qr", q, to_html(a), ctx_title, True)
            continue
        if "relances" in sub.lower() or "questions pièges" in sub.lower():
            m = re.match(r"^(«.+?»)\s*—\s*", raw)
            if not m:
                continue
            q = clean(m.group(1).strip("«» "))
            a = split_at(segs, m.end(), 0)[1]
            a = re.sub(r"^Réponse attendue\s*:\s*", "", to_html(a))
            add("jury", q, a, ctx_title, True)
            continue

        # Dossier : jurisprudence, doctrine, textes et notions en tête de paragraphe.
        first = segs[0][1] if segs else ""
        sp = lead_split(segs, (" : ", " — "))
        if not sp:
            continue
        lead = plain(sp[0])
        if len(lead) > 160 or len(plain(sp[1])) < 25:
            continue
        section = re.sub(r"^\d+(\.\d+)*\.\s*", "", sub)
        ctx = f"{ctx_title} · {section}" if section else ctx_title
        if first == "j":
            dk = decision_key(lead)
            if dk and dk in seen_dec:
                continue
            if dk:
                seen_dec.add(dk)
            add("ref", f"Portée de : {lead}", to_html(sp[1]), ctx)
        elif first == "d":
            add("doct", f"Doctrine — {lead} : que retenir ?", to_html(sp[1]), ctx)
        elif first == "t":
            add("texte", f"Que retenir de : {lead} ?", to_html(sp[1]), ctx)
        elif first == "b" and len(lead) <= 80:
            if lead.lower() in seen_terms:
                continue
            seen_terms.add(lead.lower())
            if re.match(r"^(Accroche|Fait saillant|Apport|Visa|Exemple)", lead):
                continue
            q = ask_term(lead) if "Définitions" in sub else ask_notion(lead)
            add("def" if "Définitions" in sub else "notion", q, to_html(sp[1]), ctx)

    return {"n": n, "title": title, "memo": memo, "sujets": fi["sujets"], "cards": cards}


def ask_term(term):
    # Un intitulé qui forme déjà une phrase se complète ; un mot se définit.
    if len(term.split()) > 6:
        return f"Complétez : « {term} … »"
    return f"Définition : {term}"


def ask_notion(lead):
    lead = lead.rstrip(" :;,.")
    if lead.endswith("?"):
        return lead
    if len(lead.split()) > 6:
        return f"Complétez : « {lead} … »"
    return f"{lead} : que faut-il savoir ?"


def plain_text(h):
    return re.sub(r"<[^>]+>", "", h)


# ---------- Banque de questions du jury (PDF) ----------
def parse_jury():
    path = os.path.join(SRC, "Q&A", "GrandOral_Banque_questions_jury.pdf")
    chapters, cur_ch, cur_sec, cur = [], None, None, None
    blue = (0.117647, 0.352941, 0.556863)

    def flush():
        nonlocal cur
        if cur and cur["q"] and cur["a"]:
            cur_sec["cards"].append({
                "id": f"jury-{cur['num']}", "k": "jury", "num": cur["num"], "lv": cur["lv"], "who": cur["who"],
                "q": clean(" ".join(cur["q"])),
                "a": html.escape(clean(" ".join(cur["a"])), quote=False),
                "cite": clean(" ".join(cur["cite"])).removeprefix("À citer :").strip(),
            })
        cur = None

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[5:]:
            for line in page.extract_text_lines(return_chars=True):
                text = clean(line["text"])
                chars = line["chars"]
                if not chars or not text:
                    continue
                size = round(chars[0]["size"], 1)
                bold = "Bold" in chars[0]["fontname"]
                color = tuple(round(x, 6) for x in (chars[0].get("non_stroking_color") or ()))
                if size <= 8.0 or re.fullmatch(r"\d+", text):
                    continue  # en-têtes et numéros de page
                if size >= 19:
                    flush()
                    cur_ch = {"title": text, "sections": []}
                    chapters.append(cur_ch)
                    cur_sec = None
                    continue
                m = re.match(r"^Q(\d+) ● (Base|Approfondi|Piège) · (.+)$", text)
                if m:
                    flush()
                    if cur_sec is None:
                        cur_sec = {"title": cur_ch["title"], "cards": []}
                        cur_ch["sections"].append(cur_sec)
                    cur = {"num": int(m.group(1)), "lv": m.group(2), "who": m.group(3), "q": [], "a": [], "cite": [], "in": "q"}
                    continue
                if size == 11.5 and bold:
                    flush()
                    cur_sec = {"title": text, "cards": []}
                    cur_ch["sections"].append(cur_sec)
                    continue
                if cur is None:
                    continue
                if size >= 10.4 and bold and cur["in"] == "q":
                    cur["q"].append(text)
                elif color == blue or text.startswith("À citer"):
                    cur["in"] = "cite"
                    cur["cite"].append(text)
                elif cur["in"] != "cite" and 9.5 <= size < 10.4:
                    cur["in"] = "a"
                    cur["a"].append(text)
                elif cur["in"] == "cite":
                    cur["cite"].append(text)
        flush()
    for ch in chapters:
        ch["sections"] = [s for s in ch["sections"] if s["cards"]]
    return [ch for ch in chapters if ch["sections"]]


# ---------- 500 questions PréBarreau (PDF) + mises à jour d'octobre 2026 ----------
def parse_updates():
    path = glob.glob(os.path.join(SRC, "Q&A", "Grand Oral - MAJ*.docx"))[0]
    txt = subprocess.run(["pandoc", path, "-t", "plain", "--wrap=none"], capture_output=True, text=True, check=True).stdout
    ups, cur = {}, None
    for line in txt.splitlines():
        m = re.match(r"^\s*([⟳✓]) Q(\d+)\. (.+)$", line)
        if m:
            cur = int(m.group(2))
            ups[cur] = {"changed": m.group(1) == "⟳", "q": clean(m.group(3)), "a": []}
            continue
        if cur and line.startswith("  ") and line.strip() and not re.match(r"^\s*[①-⑩]", line):
            ups[cur]["a"].append(clean(line))
        elif not line.strip():
            continue
        else:
            cur = None
    return ups


def parse_q500(updates):
    path = os.path.join(SRC, "PréBarreau", "500 Q° 2025.pdf")
    txt = subprocess.run(["pdftotext", path, "-"], capture_output=True, text=True, check=True).stdout
    lines = txt.splitlines()
    start = next(i for i, l in enumerate(lines) if l.strip() == "I- CULTURE JURIDIQUE GENERALE" and i > 500)
    sections, sec, cur, expect = [], None, None, 1
    part, pending = "", False

    def flush():
        if cur and cur["a"]:
            a = "\n".join(cur["a"])
            a = re.sub(r"\n•\n+", "\n• ", a)
            a = re.sub(r"(?<![.:;!?»)])\n(?![•\-–\d])", " ", a)
            a = re.sub(r"\n{2,}", "\n", a).strip()
            sec["cards"].append({"id": f"q500-{cur['num']}", "k": "q500", "num": cur["num"], "q": clean(cur["q"]), "a": a})

    for raw in lines[start:]:
        l = raw.strip()
        if not l or re.fullmatch(r"\d{1,3}", l):
            continue
        # Numéro seul sur sa ligne (« 73. ») : l'intitulé suit.
        if re.fullmatch(rf"{expect}\.", l):
            pending = True
            continue
        if pending:
            l, pending = f"{expect}. {l}", False
        m = re.match(r"^(\d{1,3})\.\s+(.+)$", l)
        if m and int(m.group(1)) == expect:
            flush()
            cur = {"num": expect, "q": m.group(2), "a": [], "open": not m.group(2).rstrip().endswith(("?", "»"))}
            if sec is None:
                sec = {"title": part, "cards": []}
                sections.append(sec)
            expect += 1
            continue
        if re.match(r"^[IVX]+-\s", l) and l.upper() == l:
            flush()
            cur = None
            part = l.split("- ", 1)[1].capitalize()
            sec = None
            continue
        if l.upper() == l and len(l) >= 3 and not re.search(r"\d", l) and len(l) < 70:
            flush()
            cur = None
            sec = {"title": f"{part} · {l.capitalize()}" if part else l.capitalize(), "cards": []}
            sections.append(sec)
            continue
        if cur is None:
            continue
        # Question sur deux lignes : la suite de l'intitulé vient avant la réponse.
        if cur["open"] and not cur["a"] and len(l) < 90 and (l[0].islower() or l.endswith("?")):
            cur["q"] += " " + l
            cur["open"] = False
            continue
        cur["open"] = False
        if not re.fullmatch(r"\d{1,2}\.", l):  # numéros de liste orphelins
            cur["a"].append(l)
    flush()

    # Question perdue dans la mise en page du PDF mais reprise dans la mise à jour : on la rétablit.
    have = {c["num"] for s in sections for c in s["cards"]}
    for num, u in sorted(updates.items()):
        if num not in have and u["a"]:
            s = next(s for s in sections if any(c["num"] == num - 1 for c in s["cards"]))
            i = next(i for i, c in enumerate(s["cards"]) if c["num"] == num - 1)
            s["cards"].insert(i + 1, {"id": f"q500-{num}", "k": "q500", "num": num, "q": u["q"], "a": ""})

    for s in sections:
        for c in s["cards"]:
            u = updates.get(c["num"])
            if u and u["a"]:
                c["upd"] = "oct. 2026"
                c["a"] = "\n".join(u["a"])
                if u["changed"]:
                    c["q"] = u["q"]
    return [s for s in sections if s["cards"]]


def main():
    os.makedirs(OUT, exist_ok=True)
    parts, titles = course_map()
    raw = parse_fiches(titles)
    fiches = {n: build_fiche(raw[n]) for n in sorted(raw)}
    missing = [n for p in parts for n in p["fiches"] if n not in fiches]
    if missing:
        sys.exit(f"Fiches absentes des documents : {missing}")
    cours = {"parts": parts, "fiches": [fiches[n] for n in sorted(fiches)]}
    jury = parse_jury()
    q500 = parse_q500(parse_updates())

    ids = [c["id"] for f in cours["fiches"] for c in f["cards"]]
    ids += [c["id"] for ch in jury for s in ch["sections"] for c in s["cards"]]
    ids += [c["id"] for s in q500 for c in s["cards"]]
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        sys.exit(f"Identifiants en double : {sorted(dup)[:10]}")

    for name, data in (("cours", cours), ("jury", {"chapters": jury}), ("q500", {"sections": q500})):
        with open(os.path.join(OUT, f"{name}.json"), "w", encoding="utf-8") as fh:
            json.dump(data, fh, ensure_ascii=False, separators=(",", ":"))

    nc = sum(len(f["cards"]) for f in cours["fiches"])
    nj = sum(len(s["cards"]) for ch in jury for s in ch["sections"])
    nq = sum(len(s["cards"]) for s in q500)
    print(f"Cours : {len(cours['fiches'])} fiches, {nc} cartes")
    print(f"Jury : {len(jury)} chapitres, {nj} questions")
    print(f"500 questions : {len(q500)} sections, {nq} questions")


if __name__ == "__main__":
    main()
